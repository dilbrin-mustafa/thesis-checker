"""Synthetic thesis fixtures for the three parsers.

Real theses are the Phase 0 corpus (``research/corpus/documents/``, not in git
for confidentiality/licence reasons). These builders exist so the parsers are
tested and measured on *deterministic, version-controlled* documents that
exercise the shapes that actually appear in PG theses: bilingual abstracts, a
ToC, numbered headings, numeric **and** author-year citations, figure/table
captions with cross-references, code, quotes, footnotes and a bibliography.

Every builder returns an :class:`Expected` annotation — the ground truth the
ingestion accuracy report scores against.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class Expected:
    """Hand-written ground truth for one synthetic document."""

    source_format: str
    language: str = "pl"
    headings: list[tuple[str, int]] = field(default_factory=list)
    sections: list[str] = field(default_factory=list)
    citations: list[str] = field(default_factory=list)
    references: list[dict[str, Any]] = field(default_factory=list)
    figures: int = 0
    tables: int = 0
    code_blocks: int = 0
    excluded_kinds: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def build_docx(path: Path) -> Expected:
    """A miniature PL thesis with the full PG front matter, as a .docx."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph("POLITECHNIKA GDAŃSKA", style="Title")
    doc.add_paragraph("Wydział Elektroniki, Telekomunikacji i Informatyki")
    doc.add_paragraph("Oświadczenie: niniejszą pracę wykonałem samodzielnie.")

    doc.add_heading("Streszczenie", level=1)
    doc.add_paragraph(
        "Praca przedstawia system wspomagający wstępną kontrolę jakości prac "
        "dyplomowych. Omówiono architekturę systemu oraz wyniki eksperymentów."
    )
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This thesis presents a system that supports the pre-submission quality "
        "check of engineering theses. The architecture and the experimental "
        "results are discussed in detail."
    )
    doc.add_heading("Spis treści", level=1)
    doc.add_paragraph("1. Wstęp\t5")
    doc.add_paragraph("2. Metodyka badań\t9")

    doc.add_heading("Wstęp", level=1)
    doc.add_paragraph(
        "Niniejsza praca opisuje podejście do automatycznej weryfikacji tekstów "
        "naukowych, zaproponowane pierwotnie w [1] oraz rozwinięte w (Kowalski, 2020). "
        "Prof. Nowak zauważył, że rys. 1.1 ilustruje kluczową zależność."
    )
    doc.add_heading("Metodyka badań", level=1)
    doc.add_paragraph(
        "Zastosowano metody opisane w [2, 3] oraz w [4-6]. Zestawienie wyników "
        "przedstawiono w tab. 2.1."
    )
    doc.add_paragraph("Rys. 1.1. Zależność dokładności od długości tekstu", style="Caption")
    doc.add_paragraph("Tab. 2.1. Zestawienie wyników dla trzech formatów", style="Caption")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Format"
    table.cell(0, 1).text = "F1"
    table.cell(1, 0).text = "DOCX"
    table.cell(1, 1).text = "0,93"

    code = doc.add_paragraph("def main():\n    return 0")
    code.style = doc.styles["Normal"]
    code.runs[0].font.name = "Consolas"

    quote = doc.add_paragraph(
        "Cytowany fragment pracy innego autora, przytoczony dosłownie w oryginale."
    )
    quote.style = doc.styles["Quote"] if "Quote" in [s.name for s in doc.styles] else None

    doc.add_paragraph("Kolejne punkty wyliczenia:", style="List Paragraph")
    for item in ("pierwszy punkt wyliczenia", "drugi punkt wyliczenia"):
        doc.add_paragraph(item, style="List Bullet")

    body = doc.add_paragraph("Akapit z przypisem dolnym zawierającym dodatkowe wyjaśnienie.")
    body.add_run("").font.name = "Times New Roman"

    doc.add_heading("Podsumowanie", level=1)
    doc.add_paragraph("Praca pokazuje, że automatyczna kontrola jest wykonalna i użyteczna.")

    doc.add_heading("Bibliografia", level=1)
    doc.add_paragraph(
        "[1] Kowalski J.: Metody analizy tekstu naukowego. Wydawnictwo PG, Gdańsk 2020."
    )
    doc.add_paragraph(
        "[2] Smith J., Jones A.: Neural detection of generated text. Journal of AI, "
        "vol. 12, 2019, s. 3-15. doi:10.1000/xyz123"
    )
    doc.add_paragraph(
        "[3] Nowak A., Wiśniewska M.: Wykrywanie plagiatów w języku polskim. "
        "Zeszyty Naukowe PG, 2021."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    del quote

    return Expected(
        source_format="docx",
        language="pl",
        headings=[
            # the Title-styled university line is a real heading in the source
            ("POLITECHNIKA GDAŃSKA", 1),
            ("Streszczenie", 1),
            ("Abstract", 1),
            ("Spis treści", 1),
            ("Wstęp", 1),
            ("Metodyka badań", 1),
            ("Podsumowanie", 1),
            ("Bibliografia", 1),
        ],
        sections=[
            "title_page",
            "abstract_pl",
            "abstract_en",
            "toc",
            "introduction",
            "methodology",
            "summary",
            "bibliography",
        ],
        citations=["[1]", "(Kowalski, 2020)", "[2, 3]", "[4-6]"],
        references=[
            {"year": 2020, "first_author": "Kowalski"},
            {"year": 2019, "doi": "10.1000/xyz123"},
            {"year": 2021},
        ],
        figures=1,
        tables=1,
        code_blocks=1,
        excluded_kinds=["front_matter", "toc", "bibliography", "code", "quote"],
    )


# --------------------------------------------------------------------------- #
# LaTeX
# --------------------------------------------------------------------------- #
MAIN_TEX = r"""\documentclass[12pt,a4paper]{report}
\usepackage[utf8]{inputenc}
\usepackage{polski}
\usepackage{graphicx}
\usepackage{listings}
\usepackage{amsmath}

\title{Automatyczna kontrola jakości prac dyplomowych}
\author{Jan Kowalski}

\begin{document}
\maketitle

\begin{abstract}
Praca przedstawia system wspomagający wstępną kontrolę jakości prac dyplomowych.
Omówiono architekturę systemu oraz wyniki eksperymentów.
\end{abstract}

\tableofcontents

\chapter{Wstęp}
\label{chap:wstep}
Niniejsza praca opisuje podejście do automatycznej weryfikacji tekstów
naukowych~\cite{kowalski2020} oraz metody zaproponowane w \cite{smith2019,nowak2021}.
Prof. Nowak zauważył, że rysunek~\ref{fig:dokladnosc} ilustruje kluczową zależność.

\input{chapters/metodyka}

\chapter{Podsumowanie}
Praca pokazuje, że automatyczna kontrola jest wykonalna i użyteczna
\cite{kowalski2020}.

\begin{thebibliography}{9}
\bibitem{kowalski2020} Kowalski J.: Metody analizy tekstu naukowego. Wydawnictwo PG, Gdańsk 2020.
\bibitem{smith2019} Smith J., Jones A.: Neural detection of generated text.
  Journal of AI, vol. 12, 2019, s. 3--15.
\bibitem{nowak2021} Nowak A., Wiśniewska M.: Wykrywanie plagiatów w języku
  polskim. Zeszyty Naukowe PG, 2021.
\end{thebibliography}

\end{document}
"""

METODYKA_TEX = r"""\chapter{Metodyka badań}
Zastosowano metody opisane w \cite{smith2019} oraz w \cite{nowak2021}.
Zestawienie wyników przedstawiono w tabeli~\ref{tab:wyniki}.

\begin{figure}[ht]
  \centering
  \includegraphics[width=0.8\textwidth]{img/dokladnosc.png}
  \caption{Zależność dokładności od długości tekstu}
  \label{fig:dokladnosc}
\end{figure}

\begin{table}[ht]
  \centering
  \caption{Zestawienie wyników dla trzech formatów}
  \label{tab:wyniki}
  \begin{tabular}{lc}
    Format & F1 \\
    DOCX & 0,93 \\
  \end{tabular}
\end{table}

\begin{equation}
  E = mc^2
  \label{eq:energia}
\end{equation}

\begin{itemize}
  \item pierwszy punkt wyliczenia
  \item drugi punkt wyliczenia
\end{itemize}

\begin{lstlisting}
def main():
    return 0
\end{lstlisting}

\begin{quote}
Cytowany fragment pracy innego autora, przytoczony dosłownie w oryginale.
\end{quote}

Akapit z przypisem\footnote{Treść przypisu dolnego z dodatkowym wyjaśnieniem.} w środku.
"""


def write_latex_project(directory: Path) -> Expected:
    """A multi-file LaTeX project (``\\input`` + BibTeX-free thebibliography)."""
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "main.tex").write_text(MAIN_TEX, encoding="utf-8")
    (directory / "chapters").mkdir(exist_ok=True)
    (directory / "chapters" / "metodyka.tex").write_text(METODYKA_TEX, encoding="utf-8")
    return Expected(
        source_format="latex",
        language="pl",
        headings=[
            ("Abstract", 1),
            ("Wstęp", 1),
            ("Metodyka badań", 1),
            ("Podsumowanie", 1),
            # synthesised from \begin{thebibliography} — the IR needs the heading
            ("Bibliografia", 1),
        ],
        sections=["abstract_pl", "introduction", "methodology", "summary", "bibliography"],
        # every occurrence, not the unique set: \cite{a,b} counts twice
        citations=[
            "kowalski2020",
            "smith2019",
            "nowak2021",
            "smith2019",
            "nowak2021",
            "kowalski2020",
        ],
        references=[{"year": 2020}, {"year": 2019}, {"year": 2021}],
        figures=1,
        tables=1,
        code_blocks=1,
        # front_matter covers the abstract (excluded from similarity/AIGT)
        excluded_kinds=["front_matter", "bibliography", "code", "quote"],
    )


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def build_pdf(path: Path) -> Expected:
    """A three-page PDF with an embedded ToC, mixed font sizes and a table."""
    import pymupdf

    doc = pymupdf.open()
    for _ in range(3):
        doc.new_page(width=595, height=842)
    # Page handles are re-fetched: pymupdf invalidates them when pages are added.
    pages = [doc[0], doc[1], doc[2]]

    # page 1: front matter + ToC
    pages[0].insert_htmlbox(
        pymupdf.Rect(50, 50, 545, 300),
        "<p style='font-size:20pt;font-weight:bold'>POLITECHNIKA GDAŃSKA</p>"
        "<p style='font-size:11pt'>Wydział Elektroniki, Telekomunikacji i Informatyki</p>"
        "<p style='font-size:16pt;font-weight:bold'>Streszczenie</p>"
        "<p style='font-size:12pt'>Praca przedstawia system wspomagający wstępną kontrolę "
        "jakości prac dyplomowych oraz wyniki eksperymentów.</p>"
        "<p style='font-size:16pt;font-weight:bold'>Spis treści</p>"
        "<p style='font-size:12pt'>1. Wstęp ......... 2</p>"
        "<p style='font-size:12pt'>2. Metodyka badań ......... 3</p>",
    )

    # page 2: introduction
    pages[1].insert_htmlbox(
        pymupdf.Rect(50, 50, 545, 400),
        "<p style='font-size:16pt;font-weight:bold'>1. Wstęp</p>"
        "<p style='font-size:12pt'>Niniejsza praca opisuje podejście do automatycznej "
        "weryfikacji tekstów naukowych zaproponowane w [1] oraz rozwinięte w "
        "(Kowalski, 2020). Prof. Nowak zauważył, że rys. 1.1 ilustruje zależność.</p>"
        "<p style='font-size:10pt'>Rys. 1.1. Zależność dokładności od długości tekstu</p>",
    )

    # page 3: methods + a real ruled table + bibliography
    pages[2].insert_htmlbox(
        pymupdf.Rect(50, 50, 545, 200),
        "<p style='font-size:16pt;font-weight:bold'>2. Metodyka badań</p>"
        "<p style='font-size:12pt'>Zastosowano metody opisane w [2, 3] oraz w [4-6]. "
        "Wyniki przedstawiono w tab. 2.1.</p>"
        "<p style='font-size:10pt'>Tab. 2.1. Zestawienie wyników dla trzech formatów</p>",
    )
    _draw_table(pages[2], x=50, y=250, rows=["Format  F1", "DOCX  0,93"])
    pages[2].insert_htmlbox(
        pymupdf.Rect(50, 400, 545, 600),
        "<p style='font-size:16pt;font-weight:bold'>Bibliografia</p>"
        "<p style='font-size:11pt'>[1] Kowalski J.: Metody analizy tekstu naukowego. "
        "Wydawnictwo PG, Gdańsk 2020.</p>"
        "<p style='font-size:11pt'>[2] Smith J., Jones A.: Neural detection of generated "
        "text. Journal of AI, vol. 12, 2019, s. 3-15.</p>",
    )

    doc.set_toc([[1, "Streszczenie", 1], [1, "1. Wstęp", 2], [1, "2. Metodyka badań", 3]])
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()

    return Expected(
        source_format="pdf",
        language="pl",
        headings=[
            ("POLITECHNIKA GDAŃSKA", 1),
            ("Streszczenie", 1),
            ("Spis treści", 1),
            ("Wstęp", 1),
            ("Metodyka badań", 1),
            ("Bibliografia", 1),
        ],
        sections=[
            "title_page",
            "abstract_pl",
            "toc",
            "introduction",
            "methodology",
            "bibliography",
        ],
        citations=["[1]", "(Kowalski, 2020)", "[2, 3]", "[4-6]"],
        references=[{"year": 2020, "first_author": "Kowalski"}, {"year": 2019}],
        figures=1,
        tables=1,
        code_blocks=0,
        excluded_kinds=["front_matter", "toc", "bibliography"],
    )


def _draw_table(page: Any, *, x: float, y: float, rows: list[str]) -> None:
    """Draw a ruled table so ``find_tables`` has real lines to detect."""
    import pymupdf

    row_h, col_w = 22.0, 200.0
    for row_index, row in enumerate(rows):
        top = y + row_index * row_h
        page.draw_line(pymupdf.Point(x, top), pymupdf.Point(x + col_w, top))
        page.insert_text((x + 6, top + 15), row, fontname="helv", fontsize=10)
    bottom = y + len(rows) * row_h
    page.draw_line(pymupdf.Point(x, bottom), pymupdf.Point(x + col_w, bottom))
    for offset in (0.0, col_w / 2, col_w):
        page.draw_line(pymupdf.Point(x + offset, y), pymupdf.Point(x + offset, bottom))


# --------------------------------------------------------------------------- #
# Degraded variants — these exercise the fallback paths, and they are the
# fixtures that keep the accuracy report honest.
# --------------------------------------------------------------------------- #
def build_docx_manual(path: Path) -> Expected:
    """DOCX with *no* Heading styles: headings are hand-formatted (bold + bigger).

    This is the case a template-based parser silently fails on, so the expected
    annotation deliberately includes headings that only a formatting heuristic
    can find.
    """
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def manual_heading(text: str, size: int) -> None:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        run.font.name = "Calibri"

    def body(text: str) -> None:
        para = doc.add_paragraph(text)
        para.runs[0].font.name = "Calibri"
        para.runs[0].font.size = Pt(11)

    manual_heading("Wstęp", 16)
    body(
        "Praca opisuje podejście do automatycznej weryfikacji tekstów, "
        "zaproponowane w [1] oraz rozwinięte później w (Kowalski, 2020)."
    )
    manual_heading("Metodyka badań", 16)
    body("Zastosowano metody opisane w [2] oraz porównano je z podejściem z [3].")
    manual_heading("Bibliografia", 16)
    body("[1] Kowalski J.: Metody analizy tekstu. Wydawnictwo PG, Gdańsk 2020.")
    body("[2] Smith J.: Neural detection. Journal of AI, 2019, s. 3-15.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return Expected(
        source_format="docx",
        language="pl",
        headings=[("Wstęp", 1), ("Metodyka badań", 1), ("Bibliografia", 1)],
        sections=["introduction", "methodology", "bibliography"],
        citations=["[1]", "(Kowalski, 2020)", "[2]", "[3]"],
        references=[{"year": 2020, "first_author": "Kowalski"}, {"year": 2019}],
        figures=0,
        tables=0,
        code_blocks=0,
        excluded_kinds=["bibliography"],
    )


def build_pdf_no_toc(path: Path) -> Expected:
    """PDF with no embedded ToC: heading levels come from font-size clustering."""
    import pymupdf

    doc = pymupdf.open()
    for _ in range(2):
        doc.new_page(width=595, height=842)
    pages = [doc[0], doc[1]]

    pages[0].insert_htmlbox(
        pymupdf.Rect(50, 50, 545, 400),
        "<p style='font-size:14pt;font-weight:bold'>Wstęp</p>"
        "<p style='font-size:11pt'>Niniejsza praca opisuje podejście do automatycznej "
        "weryfikacji tekstów naukowych zaproponowane w [1]. Prof. Nowak zauważył, "
        "że rys. 1.1 ilustruje zależność dokładności od długości tekstu.</p>"
        "<p style='font-size:11pt'>Akapit bez cytatu, ale z odwołaniem do tab. 1.1 "
        "zawierającej zestawienie wyników eksperymentów.</p>"
        "<p style='font-size:9pt'>Rys. 1.1. Zależność dokładności od długości tekstu</p>"
        "<p style='font-size:9pt'>Tab. 1.1. Zestawienie wyników eksperymentów</p>",
    )
    pages[1].insert_htmlbox(
        pymupdf.Rect(50, 50, 545, 400),
        "<p style='font-size:14pt;font-weight:bold'>Bibliografia</p>"
        "<p style='font-size:11pt'>[1] Kowalski J.: Metody analizy tekstu naukowego. "
        "Wydawnictwo PG, Gdańsk 2020.</p>",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return Expected(
        source_format="pdf",
        language="pl",
        headings=[("Wstęp", 1), ("Bibliografia", 1)],
        sections=["introduction", "bibliography"],
        citations=["[1]"],
        references=[{"year": 2020, "first_author": "Kowalski"}],
        figures=1,
        tables=1,
        code_blocks=0,
        excluded_kinds=["bibliography"],
    )


ARTICLE_TEX = r"""\documentclass[12pt]{article}
\usepackage[utf8]{inputenc}
\usepackage{polski}

\title{Analiza sygnałów}
\author{Anna Nowak}

\begin{document}
\maketitle

\section{Wstęp}
Praca opisuje metody analizy sygnałów zastosowane w \cite{kowalski2020}.

\subsection{Zakres pracy}
Zakres obejmuje sygnały jednowymiarowe opisane w \cite{smith2019}.

\section{Wyniki}
Wyniki przedstawiono poniżej; porównanie z \cite{kowalski2020} wypada korzystnie.

\bibliographystyle{plain}
\bibliography{refs}

\end{document}
"""

REFS_BIB = r"""@book{kowalski2020,
  author    = {Kowalski, Jan},
  title     = {Metody analizy tekstu naukowego},
  publisher = {Wydawnictwo PG},
  year      = {2020},
  address   = {Gdańsk}
}

@article{smith2019,
  author  = {Smith, John and Jones, Alice},
  title   = {Neural detection of generated text},
  journal = {Journal of AI},
  volume  = {12},
  pages   = {3--15},
  year    = {2019},
  doi     = {10.1000/xyz123}
}
"""


def write_latex_article(directory: Path) -> Expected:
    """``article``-class project with a real ``.bib`` file (bibtexparser path)."""
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "main.tex").write_text(ARTICLE_TEX, encoding="utf-8")
    (directory / "refs.bib").write_text(REFS_BIB, encoding="utf-8")
    return Expected(
        source_format="latex",
        language="pl",
        headings=[("Wstęp", 1), ("Zakres pracy", 2), ("Wyniki", 1)],
        # \bibliography{} points at a .bib file: the bibliography never appears
        # in the document text, so there is neither a section nor anything to
        # exclude. References still reach the IR, straight from the .bib.
        sections=["introduction", "results"],
        citations=["kowalski2020", "smith2019", "kowalski2020"],
        references=[
            {"year": 2020, "first_author": "Kowalski"},
            {"year": 2019, "doi": "10.1000/xyz123"},
        ],
        figures=0,
        tables=0,
        code_blocks=0,
        excluded_kinds=[],
    )
